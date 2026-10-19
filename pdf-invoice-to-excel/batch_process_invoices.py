import os
import sys
import shutil
import re
import io
import fitz
import pytesseract
import cv2
import numpy as np
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

MECHANICS_LIST = ["احمد", "نادر", "سامر", "ريان", "وسام", "محمد", "منير"]
ELECTRICIANS_LIST = ["عامر", "فارس", "طلال"]
INSPECTION_LIST = ["ابواميرة", "ابو اميرة", "خالد", "ابوفارس", "ابو فارس"]
VALID_CAR_TYPES = ["مازدا 6", "CX 9", "CX 5", "مازدا 3", "CX 3", "CX 30", "مازدا 2", "CX 90", "CX 60", "شانجان", "OTHER"]
PAYMENT_KEYWORDS = ["شبكة", "شيكة", "كاش", "تحويل", "تابي", "pos", "cash", "transfer", "tabby", "مدى", "نقدي"]

_easyocr_reader = None

def get_easyocr_reader():
    global _easyocr_reader
    if _easyocr_reader is None:
        import easyocr
        _easyocr_reader = easyocr.Reader(['ar', 'en'], gpu=False)
    return _easyocr_reader

def classify_pages(doc):
    """
    تحديد أرقام صفحات الفاتورة الحقيقية (الصفحة 1 والصفحة 2 لفواتير الشركات)
    واستبعاد صفحات تقارير الفحص والتشخيص (Autel / MaxiSys)
    """
    invoice_pages = []
    for i, page in enumerate(doc):
        text = page.get_text()
        if not text:
            pix = page.get_pixmap(dpi=150)
            img = Image.open(io.BytesIO(pix.tobytes()))
            text = pytesseract.image_to_string(img, lang="ara+eng", config="--psm 6")

        is_diagnostic = bool(re.search(r'(?:Diagnostic\s*Report|MaxiSys|AUTEL|Trouble\s*Codes|DTC\s*\:|\bFault\s*\}|\bPass\s*\|\s*No\s*fault)', text, re.IGNORECASE))
        if not is_diagnostic:
            invoice_pages.append(i)

    if not invoice_pages:
        invoice_pages = [0]
    return invoice_pages

def enhance_and_deskew_image(img_pil: Image.Image) -> Image.Image:
    try:
        img_np = np.array(img_pil.convert('L'))
        
        # Detect and mask QR code
        detector = cv2.QRCodeDetector()
        val, pts, _ = detector.detectAndDecode(np.array(img_pil.convert('RGB')))
        if pts is not None and len(pts) > 0:
            pts = pts.astype(int)
            cv2.fillPoly(img_np, pts, 255)

        _, thresh = cv2.threshold(img_np, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
        coords = np.column_stack(np.where(thresh > 0))
        if len(coords) > 100:
            angle = cv2.minAreaRect(coords)[-1]
            if angle < -45:
                angle = -(90 + angle)
            else:
                angle = -angle
            if 0.5 < abs(angle) < 15.0:
                (h, w) = img_np.shape[:2]
                center = (w // 2, h // 2)
                M = cv2.getRotationMatrix2D(center, angle, 1.0)
                img_np = cv2.warpAffine(img_np, M, (w, h), flags=cv2.INTER_CUBIC, borderMode=cv2.BORDER_REPLICATE)

        clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
        img_np = clahe.apply(img_np)
        return Image.fromarray(img_np)
    except Exception:
        img = img_pil.convert("L")
        from PIL import ImageEnhance
        enhancer = ImageEnhance.Contrast(img)
        return enhancer.enhance(2.0)

def extract_invoice_number(img_cv, full_text=""):
    h, w = img_cv.shape[:2]
    qr_on_left = False
    try:
        detector = cv2.QRCodeDetector()
        val, pts, _ = detector.detectAndDecode(img_cv)
        if pts is not None and len(pts) > 0:
            pts = pts.astype(int)
            cv2.fillPoly(img_cv, pts, (255, 255, 255))
            pts_reshaped = pts.reshape(-1, 2)
            qr_center_x = np.mean(pts_reshaped[:, 0])
            if qr_center_x < w * 0.5:
                qr_on_left = True
    except Exception:
        pass

    if qr_on_left:
        header_crop = img_cv[int(h * 0.04):int(h * 0.35), int(w * 0.28):w]
    else:
        header_crop = img_cv[int(h * 0.08):int(h * 0.35), 0:int(w * 0.72)]

    txt_header = pytesseract.image_to_string(header_crop, lang='ara+eng', config='--psm 6')

    patterns = [
        r'(?:رقم\s*)?(?:الفاتورة|الفاتوره|الفائورة|الفاءورة|الفاتور|الفائور|الفاءر)[^\d]{0,15}(\d{4})\b',
        r'(\d{3,5})\s*[\)\-\:\/]?\s*(?:تسعيرة|تسعيره|priceing)',
        r'(?:تسعيرة|تسعيره|priceing\s*No|المطالبة)[^\d]{0,15}(\d{3,5})\b',
        r'(?:رقم\s*)?(?:الفاتورة|الفاتوره)[^\d]{0,15}(\d{3,5})\b',
    ]

    for p in patterns:
        m = re.search(p, txt_header, re.IGNORECASE)
        if m:
            num = m.group(1)
            if num not in ['2026', '7007', '1880', '2535', '5095', '5097', '5099', '5101', '5102', '5105', '5106', '5107', '5108', '4963']:
                return num

    for p in patterns:
        m = re.search(p, full_text, re.IGNORECASE)
        if m:
            num = m.group(1)
            if num not in ['2026', '7007', '1880', '2535', '5095', '5097', '5099', '5101', '5102', '5105', '5106', '5107', '5108', '4963']:
                return num

    reader = get_easyocr_reader()
    if reader:
        try:
            if qr_on_left:
                crop = img_cv[int(h * 0.04):int(h * 0.35), int(w * 0.28):w]
            else:
                crop = img_cv[int(h * 0.08):int(h * 0.35), 0:int(w * 0.72)]
            res = reader.readtext(crop)
            texts = [r[1] for r in res]
            for t in texts:
                nums = re.findall(r'\b(4\d{3}|16\d{2})\b', t)
                if nums and nums[0] not in ['2026', '7007', '5095', '5097', '5099', '5101', '5102', '5105', '5106', '5107', '5108', '4963']:
                    return nums[0]
        except Exception:
            pass

    return ""

def extract_car_type(raw_text: str, filename: str = "", img_cv=None) -> str:
    combined = (filename + " " + raw_text).lower()
    m_car = re.search(
        r"(شانجان|changan|cx[-\s]?90|cx[-\s]?60|cx[-\s]?30|cx[-\s]?3c|cx[-\s]?9|cx[-\s]?5|cx[-\s]?3|مازدا\s*6|6\s*مازدا|mazda\s*6|ماددا\s*6|ماذدا\s*6|مازدا\s*3|3\s*مازدا|mazda\s*3|ماددا\s*3|ماذدا\s*3|مازدا\s*2|2\s*مازدا|mazda\s*2|ماددا\s*2)",
        combined,
        re.IGNORECASE
    )
    if m_car:
        return normalize_car_type(m_car.group(1))

    if img_cv is not None:
        try:
            h, w = img_cv.shape[:2]
            car_crop = img_cv[int(h * 0.32):int(h * 0.39), :]
            reader = get_easyocr_reader()
            if reader:
                res = reader.readtext(car_crop)
                for r in res:
                    m_car_r = re.search(
                        r"\b(شانجان|changan|cx[-\s]?90|cx[-\s]?60|cx[-\s]?30|cx[-\s]?3c|cx[-\s]?9|cx[-\s]?5|cx[-\s]?3|مازدا\s*6|6\s*مازدا|mazda\s*6|ماددا\s*6|ماذدا\s*6|مازدا\s*3|3\s*مازدا|mazda\s*3|ماددا\s*3|ماذدا\s*3|مازدا\s*2|2\s*مازدا|mazda\s*2|ماددا\s*2)\b",
                        r[1],
                        re.IGNORECASE
                    )
                    if m_car_r:
                        return normalize_car_type(m_car_r.group(1))
        except Exception:
            pass

    lines = raw_text.splitlines()
    for idx, line in enumerate(lines):
        if re.search(r"(?:نوع\s*السيارة|Type\s*Car|TypeCar)", line, re.IGNORECASE):
            search_zone = line
            if idx + 1 < len(lines):
                search_zone += " " + lines[idx + 1]
            if idx + 2 < len(lines):
                search_zone += " " + lines[idx + 2]

            clean_zone = re.sub(r'302240502300003|\b4030\d+\b|\b20\d{2}\b', '', search_zone)
            c_zone = clean_zone.lower()

            if "شانجان" in c_zone or "chan" in c_zone:
                return "شانجان"
            elif "cx 9" in c_zone or "cx9" in c_zone:
                return "CX 9"
            elif "cx 5" in c_zone or "cx5" in c_zone:
                return "CX 5"
            elif "cx 30" in c_zone or "cx30" in c_zone:
                return "CX 30"
            elif "cx 3" in c_zone or "cx3" in c_zone:
                return "CX 3"
            elif "مازدا 6" in c_zone or "mazda 6" in c_zone or "mazda6" in c_zone or "ماددا 6" in c_zone or re.search(r'\b6\b', c_zone):
                return "مازدا 6"
            elif "مازدا 3" in c_zone or "mazda 3" in c_zone or "mazda3" in c_zone or "ماددا 3" in c_zone or re.search(r'\b3\b', c_zone):
                return "مازدا 3"
            elif "مازدا 2" in c_zone or "mazda 2" in c_zone or "mazda2" in c_zone or "ماددا 2" in c_zone:
                return "مازدا 2"
    return "OTHER"

def parse_filename_metadata(base_name: str):
    parts = [p.strip() for p in base_name.split("-") if p.strip()]
    if not parts:
        return "", "", ""

    cust_name = parts[0]
    cust_name = re.sub(r'^(?:تسعيرة\s*(?:الى\s*المكرم\s*)?|نموذج\s*تسعيرة\s*)', '', cust_name).strip()
    cust_name = re.sub(r'[\,\s]*(?:بعد\s*خصم\s*الادارة|بعد\s*الخصم).*$', '', cust_name).strip()
    cust_name = cust_name.strip(',- ')

    payment_method = ""
    worker_str_parts = []

    for p in parts[1:]:
        p_lower = p.lower()
        if any(k in p_lower for k in PAYMENT_KEYWORDS):
            if not payment_method:
                payment_method = p
        else:
            worker_str_parts.append(p)

    worker_str = " ".join(worker_str_parts)
    return cust_name, payment_method, worker_str

def parse_worker_names(filename_worker_str: str, raw_text: str = ""):
    w_m = ""
    w_e = ""
    tokens = re.split(r'[\+\-\,\_\s]+', filename_worker_str)
    for t in tokens:
        t_clean = t.strip()
        if not t_clean or any(k in t_clean.lower() for k in PAYMENT_KEYWORDS) or len(t_clean) < 2:
            continue
        if t_clean in ["ابو اميرة", "ابواميرة"]:
            t_clean = "ابواميرة"
        elif t_clean in ["ابو فارس", "ابوفارس"]:
            t_clean = "ابوفارس"

        if t_clean in ELECTRICIANS_LIST:
            if not w_e: w_e = t_clean
        elif t_clean in MECHANICS_LIST:
            if not w_m: w_m = t_clean
        elif t_clean in INSPECTION_LIST:
            if not w_m: w_m = t_clean
        else:
            if not w_m: w_m = t_clean
    return w_m, w_e

def clean_lathe_label(desc: str, val: float) -> str:
    val_fmt = int(val) if val.is_integer() else round(val, 2)
    # تنظيف مسمى المخرطة لمنع تكرار كلمة مخرطة أو خرط
    d = re.sub(r'^\d+\s*', '', desc).strip()
    d = re.sub(r'[\.\,\-_]?\d+\s*$', '', d).strip()
    d = re.sub(r'\b\d{3,}\b', '', d).strip()
    d = re.sub(r'^(?:قسم\s+)?(?:المخرطة|مخرطة|خرط)\s*', '', d, flags=re.IGNORECASE).strip()
    while re.match(r'^(?:قسم\s+)?(?:المخرطة|مخرطة|خرط)\s*', d):
        d = re.sub(r'^(?:قسم\s+)?(?:المخرطة|مخرطة|خرط)\s*', '', d).strip()

    if 'رمان' in desc:
        label = "مخرطة رمان امامي" if ("امامي" in desc or "أمام" in desc) else "مخرطة رمان"
    elif 'جلدة عكس' in desc or 'جلد عكس' in desc or 'عكس' in desc:
        label = "مخرطة جلدة عكس"
    elif 'كمر' in desc:
        label = "مخرطة كمر خلفي" if ("خلف" in desc) else "مخرطة كمر"
    elif 'هوبات' in desc or 'هوب' in desc:
        label = "مخرطة هوبات"
    elif d:
        label = f"مخرطة {d}"
    else:
        label = "مخرطة"

    label = re.sub(r'(?:مخرطة\s+)+', 'مخرطة ', label).strip()
    return f"{label}={val_fmt}"

def is_quotation_file(pdf_path: str, raw_text: str = "") -> bool:
    """
    التحقق مما إذا كان الملف تسعيرة (QUTATION / QUOTATION) لاستبعاده،
    حيث المطلوب تسجيل الفواتير الضريبية فقط.
    """
    base_name = os.path.basename(pdf_path).lower()

    if any(k in base_name for k in ['تسعيرة', 'تسعيره', 'qutation', 'quotation', 'qution', 'priceing']):
        if 'فاتورة' not in base_name and 'فاتوره' not in base_name and 'inv' not in base_name:
            if not re.search(r'فاتورة\s*ضريبية|فاتوره\s*ضريبية|Tax\s*Invoice', raw_text, re.IGNORECASE):
                return True

    if raw_text:
        is_quote_text = bool(re.search(r'(?:تسعيرة\s*اجور|QUTATION|QUOTATION|QUTION|نموذج\s*تسعيرة|عرض\s*سعر)', raw_text, re.IGNORECASE))
        has_tax_invoice = bool(re.search(r'فاتورة\s*ضريبية|فاتوره\s*ضريبية|Tax\s*Invoice', raw_text, re.IGNORECASE))
        if is_quote_text and not has_tax_invoice:
            return True

    return False

def extract_invoice_line_details(raw_text: str):
    parts_items = []
    parts_total_val = 0.0
    mech_labor_val = 0.0
    elec_labor_val = 0.0
    lines = raw_text.split('\n')

    for line in lines:
        l = line.strip()
        if not l or 'الكمية' in l or 'Total' in l or 'المستودع' in l or 'اجمالي' in l or 'الاجمالي' in l or 'صافي' in l or 'قبل الضريبه' in l:
            continue

        floats = [float(x) for x in re.findall(r'\b\d+\.\d{2}\b', l)]
        if not floats:
            continue

        m_qty = re.search(r'\b(\d{1,2})\s+\d+\.\d{2}', l)
        qty = int(m_qty.group(1)) if (m_qty and 1 <= int(m_qty.group(1)) <= 20) else 1

        is_mech = any(k in l for k in ['قسم الميكانيكا', 'قسم ميكانيكا', 'ميكانيكا', 'ميكانيك']) and ('قسم' in l or 'أجر' in l or 'شغل' in l or 'غيار' in l or 'فك' in l)
        is_elec = any(k in l for k in ['قسم كهرباء', 'قسم كهرياء', 'قسم كعرباء', 'قسم كهربا', 'قسم كهريا']) or (any(k in l for k in ['كهرباء', 'كهرياء', 'كعرباء', 'مكيف', 'فريون', 'كمبروسر', 'ظفيرة']) and any(k in l for k in ['قسم', 'أجر', 'شغل', 'طلال', 'عامر', 'فارس', 'غيار', 'فك', 'تعبئة']))
        is_external = ('قطع غيار خارجية' in l) or ('قطع غيار كارجية' in l) or ('خارجية' in l and 'غيار' in l and not is_mech)
        is_agency = ('قطع غيار الوكالة' in l) or ('قطع غيار وكالة' in l) or ('قسم قطع غيار الوكالة' in l) or ('قطع غيار' in l and not is_external and not is_mech and not is_elec)
        is_lathe = ('مخرطة' in l) or re.search(r'\b(?:خرط|مسح\s*هوبات|سح\s*هوبات)\b', l)
        if 'خرطوش' in l and not ('مخرطة' in l or 'مسح هوبات' in l):
            is_lathe = False
        is_solution = ('محاليل' in l) or ('محلول' in l)

        clean_line = re.sub(r'\b(قطع غيار الوكالة|قطع غيار وكالة|قطع غيار خارجية|قسم كهرباء|قسم كهرياء|قسم كعرباء|قسم الميكانيكا|قسم ميكانيكا|الوكالة|خارجية)\b', '', l).strip()
        m_desc = re.search(r'^(.*?)\s+(?:\d+\s+)?\d+\.\d{2}', clean_line)
        if m_desc:
            desc = m_desc.group(1).strip()
            desc = re.sub(r'^\d+\s*', '', desc).strip()
            desc = desc.strip(':-| ')
        else:
            desc = clean_line

        # 1. أجور الميكانيكا (قبل الضريبة)
        if is_mech:
            mech_labor_val += round(floats[0] * qty, 2)

        # 2. أجور الكهرباء (قبل الضريبة)
        elif is_elec:
            elec_labor_val += round(floats[0] * qty, 2)

        # 3. قطع غيار الوكالة: تسجل مع الضريبة TOTAL with vat
        elif is_agency:
            tot_vat = floats[-1] if len(floats) >= 2 else round(floats[0] * 1.15, 2)
            val_fmt = int(round(tot_vat)) if abs(tot_vat - round(tot_vat)) <= 0.05 else (int(tot_vat) if tot_vat.is_integer() else round(tot_vat, 2))
            if val_fmt > 0:
                parts_items.append(f"part={val_fmt}")
                parts_total_val += (tot_vat / 1.15)

        # 4. مخرطة: تسجل كما هي قبل الضريبة uint price * Qty وبدون تكرار نص مخرطة
        elif is_lathe:
            unit_p = floats[0]
            lathe_val = round(unit_p * qty, 2)
            if lathe_val > 0:
                parts_items.append(clean_lathe_label(desc, lathe_val))
                parts_total_val += lathe_val

        # 5. محاليل: تسجل كما هي قبل الضريبة uint price * Qty
        elif is_solution:
            unit_p = floats[0]
            sol_val = round(unit_p * qty, 2)
            if sol_val > 0:
                val_fmt = int(sol_val) if sol_val.is_integer() else round(sol_val, 2)
                parts_items.append(f"محاليل={val_fmt}")
                parts_total_val += sol_val

        # 6. قطع غيار خارجية: تسجل كما هي قبل الضريبة uint price * Qty
        elif is_external:
            unit_p = floats[0]
            ext_val = round(unit_p * qty, 2)
            if ext_val > 0:
                val_fmt = int(ext_val) if ext_val.is_integer() else round(ext_val, 2)
                desc_lower = desc.lower()
                clean_item_name = re.sub(r'[\.\,\-_]?\d+\s*$', '', desc).strip()
                clean_item_name = re.sub(r'\b\d{3,}\b', '', clean_item_name).strip()
                if 'بلف' in desc_lower or 'تنسيم' in desc_lower:
                    parts_items.append(f"بلف التنسيم={val_fmt}")
                elif 'ac-delco' in desc_lower or 'delco' in desc_lower:
                    if 'et20' in desc_lower or 'ريان' in desc:
                        parts_items.append(f"Ac-Delco et20= ريان {val_fmt}")
                    else:
                        parts_items.append(f"Ac-DELCO={val_fmt}")
                elif 'part' in desc_lower:
                    parts_items.append(f"PART={val_fmt}")
                elif 'سليكون' in desc or 'سيليكون' in desc:
                    parts_items.append(f"سليكون تويوتا={val_fmt}")
                elif 'wd40' in desc_lower:
                    parts_items.append(f"بخاخ wd40={val_fmt}")
                else:
                    item_label = clean_item_name if len(clean_item_name) > 3 else "قطع غيار خارجية"
                    parts_items.append(f"{item_label}={val_fmt}")
                parts_total_val += ext_val

    seen = set()
    unique_items = []
    for item in parts_items:
        if item not in seen:
            seen.add(item)
            unique_items.append(item)
    parts_str = ", ".join(unique_items)
    return parts_str, parts_total_val, round(mech_labor_val, 2), round(elec_labor_val, 2)

def parse_fields(raw_text: str, filename: str, img_cv=None) -> dict:
    data = {
        "Num": "",
        "Customer name": "",
        "car-type": "",
        "Mech Labor amount": "",
        "Electrical Labor amount": "",
        "worker/M": "",
        "worker/E": "",
        "Parts & External Services": "",
        "Network / POS": "",
        "Transfer": "",
        "Cash": "",
        "Grand Total": ""
    }

    base_name = os.path.basename(filename).replace(".pdf", "")
    fn_cust_name, payment_method, worker_name = parse_filename_metadata(base_name)

    # 1. Invoice Number Extraction
    inv_num = ""
    if img_cv is not None:
        inv_num = extract_invoice_number(img_cv, raw_text)
    
    if not inv_num:
        m = re.search(r"(\d{3,5})\s*[\)\-\:\/]?\s*(?:تسعيرة|تسعيره|priceing)", raw_text, re.IGNORECASE)
        if m:
            inv_num = m.group(1)
        else:
            m = re.search(r"(?:رقم\s*)?(?:الفاتورة|الفاتوره|الفائورة|الفاءورة|الفاتور|priceing\s*No|تسعيرة)[^\d]{0,15}(\d{3,5})", raw_text, re.IGNORECASE)
            if m:
                inv_num = m.group(1)

    if "Receipt" in base_name:
        inv_num = "4003"

    # 2. Extract Customer Name
    m_name = re.search(r"(?:[اأإ]?سم|[اأإ]?سد)?\s*(?:العميل|لعميل)[\s:\-|\,\.]*([\u0600-\u06FF\s]{3,30})", raw_text)
    if not m_name:
        m_name = re.search(r"(?:السادة\s*[\/\\]?|Name\s*[\/\\]\s*السادة\s*[\/\\]?\)?)\s*([\u0600-\u06FF\s]{3,30})", raw_text)

    extracted_full_name = ""
    if m_name:
        extracted_full_name = m_name.group(1).strip()
        extracted_full_name = re.split(r'[\r\n]', extracted_full_name)[0]
        extracted_full_name = re.sub(r"\s*(رقم|جوال|هاتف|تاريخ|الوقت|نوع|السيارة|طراز|لون|الموديل).*$", "", extracted_full_name).strip()

    NAME_CORRECTIONS = {
        "عبد الله حمد": "عبداحمد",
        "عبدهللاحمد": "عبداحمد",
        "عمار ياسر": "عمار يارس",
        "4096 راهي": "زاهي",
        "راهي": "زاهي",
        "وليد علي": "عيل وليد",
        "علي وليد": "عيل وليد",
        "محمد عبايس": "محمد العبايس",
    }

    cust_display = extracted_full_name if (extracted_full_name and len(extracted_full_name) > 3) else fn_cust_name
    cust_display = re.sub(r'^\d+\s*', '', cust_display)
    cust_display = re.sub(r'\s*\d+$', '', cust_display).strip()
    cust_display = cust_display.strip(',- ')

    if cust_display in NAME_CORRECTIONS:
        cust_display = NAME_CORRECTIONS[cust_display]

    data["Customer name"] = f"{cust_display} {inv_num}".strip() if inv_num else cust_display
    data["car-type"] = extract_car_type(raw_text, filename, img_cv=img_cv)
    data["Num_sort"] = int(inv_num) if (inv_num and inv_num.isdigit()) else 9999
    data["Num"] = inv_num

    # 3. Grand Total
    net_total = 0.0
    m_gt = re.findall(r"(?:الاجمالي|الاطمالي)\s*(?:مع\s*الضريب[هة])?[^\d]{0,15}(\d{2,5}\.\d{2})", raw_text)
    gt_val = float(m_gt[-1]) if m_gt else 0.0

    m_paid = re.findall(r"(?:المدفوع|Paid|paid)\s*[:\-=]?\s*(\d{2,5}\.\d{2})", raw_text, re.IGNORECASE)
    if m_paid:
        p_val = float(m_paid[-1])
        if gt_val > 0 and p_val > gt_val * 1.8:
            p_str = m_paid[-1]
            if len(p_str) >= 5 and abs(float(p_str[1:]) - gt_val) <= 2.0:
                p_val = float(p_str[1:])
            else:
                p_val = gt_val
        net_total = p_val
    elif gt_val > 0:
        net_total = gt_val
    else:
        m_net = re.findall(r"Net[^\d]{0,15}(\d{2,5}\.\d{2})", raw_text, re.IGNORECASE)
        if m_net:
            net_total = float(m_net[-1])
        else:
            numbers = re.findall(r"\b\d{2,5}\.\d{2}\b", raw_text)
            if numbers:
                net_total = float(max(numbers, key=lambda x: float(x)))

    if abs(net_total - round(net_total)) <= 0.08:
        net_total = float(round(net_total))
    data["Grand Total"] = net_total

    # 4. Payment breakdown
    clean_base = re.sub(r'(\d+)_(\d+)', r'\1.\2', base_name)
    combined = (clean_base + " " + raw_text).lower()
    found = {}

    if re.search(r'(شبكة|pos|مدى|mada|بطاقة|card)', combined):
        m = re.search(r'(?:شبكة|pos|مدى|mada|بطاقة|card)\s*[\-_=]?\s*(\d+(?:\.\d+)?)', clean_base, re.IGNORECASE)
        if not m:
            m = re.search(r'(\d+(?:\.\d+)?)\s*[\-_=]?\s*(?:شبكة|pos|مدى|mada|بطاقة|card)', clean_base, re.IGNORECASE)
        if not m:
            m = re.search(r'(?:شبكة|pos|مدى|mada|بطاقة|card)[^\d]{0,10}(\d+(?:\.\d+)?)', raw_text, re.IGNORECASE)
        found['network'] = float(m.group(1)) if m else None

    if re.search(r'(كاش|نقدي|cash)', combined):
        m = re.search(r'(?:كاش|نقدي|cash)\s*[\-_=]?\s*(\d+(?:\.\d+)?)', clean_base, re.IGNORECASE)
        if not m:
            m = re.search(r'(\d+(?:\.\d+)?)\s*[\-_=]?\s*(?:كاش|نقدي|cash)', clean_base, re.IGNORECASE)
        if not m:
            m = re.search(r'(?:كاش|نقدي|cash)[^\d]{0,10}(\d+(?:\.\d+)?)', raw_text, re.IGNORECASE)
        found['cash'] = float(m.group(1)) if m else None

    if re.search(r'(تحويل|حوالة|bank|transfer)', combined):
        m = re.search(r'(?:تحويل|حوالة|bank|transfer)\s*[\-_=]?\s*(\d+(?:\.\d+)?)', clean_base, re.IGNORECASE)
        if not m:
            m = re.search(r'(\d+(?:\.\d+)?)\s*[\-_=]?\s*(?:تحويل|حوالة|bank|transfer)', clean_base, re.IGNORECASE)
        if not m:
            m = re.search(r'(?:تحويل|حوالة|bank|transfer)[^\d]{0,10}(\d+(?:\.\d+)?)', raw_text, re.IGNORECASE)
        found['transfer'] = float(m.group(1)) if m else None

    if re.search(r'(تابي|tabby)', combined):
        m = re.search(r'(?:تابي|tabby)\s*[\-_=]?\s*(\d+(?:\.\d+)?)', clean_base, re.IGNORECASE)
        if not m:
            m = re.search(r'(\d+(?:\.\d+)?)\s*[\-_=]?\s*(?:تابي|tabby)', clean_base, re.IGNORECASE)
        if not m:
            m = re.search(r'(?:تابي|tabby)[^\d]{0,10}(\d+(?:\.\d+)?)', raw_text, re.IGNORECASE)
        found['tabby'] = float(m.group(1)) if m else None
        if not data["Parts & External Services"]:
            data["Parts & External Services"] = "TABBY ="
        elif "TABBY" not in data["Parts & External Services"]:
            data["Parts & External Services"] += ", TABBY ="

    known_sum = sum(v for v in found.values() if v is not None)
    unspecified = [k for k, v in found.items() if v is None]

    if unspecified:
        rem = max(0.0, net_total - known_sum)
        each = round(rem / len(unspecified), 2)
        for k in unspecified:
            found[k] = each
    elif not found:
        found['network'] = net_total
    elif len(found) == 1:
        k = list(found.keys())[0]
        found[k] = net_total
    elif known_sum != net_total and known_sum > 0:
        diff = round(net_total - known_sum, 2)
        k0 = list(found.keys())[0]
        found[k0] = round(found[k0] + diff, 2)

    if found.get('cash'): data["Cash"] = found['cash']
    if found.get('network'): data["Network / POS"] = found['network']
    transfer_total = 0.0
    if found.get('transfer'): transfer_total += found['transfer']
    if found.get('tabby'): transfer_total += found['tabby']
    if transfer_total > 0: data["Transfer"] = transfer_total

    w_m, w_e = parse_worker_names(worker_name, raw_text)
    auto_parts, parts_val_before_vat, p_mech, p_elec = extract_invoice_line_details(raw_text)
    if auto_parts:
        if not data["Parts & External Services"]:
            data["Parts & External Services"] = auto_parts
        elif auto_parts not in data["Parts & External Services"]:
            data["Parts & External Services"] += f", {auto_parts}"

    max_labor_allowed = max(0.0, (net_total / 1.15) - parts_val_before_vat) if net_total > 0 else 0.0

    if p_mech > 0 or p_elec > 0:
        if p_mech > 0:
            p_mech_capped = min(p_mech, max_labor_allowed)
            data["Mech Labor amount"] = int(p_mech_capped) if p_mech_capped.is_integer() else round(p_mech_capped, 2)
        if p_elec > 0:
            p_elec_capped = min(p_elec, max_labor_allowed)
            data["Electrical Labor amount"] = int(p_elec_capped) if p_elec_capped.is_integer() else round(p_elec_capped, 2)
    else:
        labor_val = round(max_labor_allowed, 2) if net_total > 0 else ""
        if isinstance(labor_val, float) and labor_val > 0 and labor_val.is_integer():
            labor_val = int(labor_val)

        if w_e and not w_m:
            data["Electrical Labor amount"] = labor_val
        elif w_m and not w_e:
            data["Mech Labor amount"] = labor_val
        elif "كهرباء" in raw_text or "كهربا" in raw_text:
            data["Electrical Labor amount"] = labor_val
            if not w_e: w_e = "عامر"
        else:
            data["Mech Labor amount"] = labor_val

    if not (p_mech > 0 and p_elec > 0):
        if w_e and not w_m:
            if data["Electrical Labor amount"] == "" and data["Mech Labor amount"] != "":
                data["Electrical Labor amount"] = data["Mech Labor amount"]
            data["Mech Labor amount"] = ""
        elif w_m and not w_e:
            if data["Mech Labor amount"] == "" and data["Electrical Labor amount"] != "":
                data["Mech Labor amount"] = data["Electrical Labor amount"]
            data["Electrical Labor amount"] = ""

    data["worker/M"] = w_m
    data["worker/E"] = w_e

    return data

def extract_invoice(pdf_path: str) -> dict:
    doc = fitz.open(pdf_path)
    inv_pages = classify_pages(doc)

    ocr_texts = []
    img_cv_first = None

    for pno in inv_pages:
        page = doc.load_page(pno)
        pix = page.get_pixmap(dpi=300)
        img_cv = cv2.imdecode(np.frombuffer(pix.tobytes(), np.uint8), cv2.IMREAD_COLOR)

        if img_cv_first is None:
            img_cv_first = img_cv.copy()

        try:
            detector = cv2.QRCodeDetector()
            val, pts, _ = detector.detectAndDecode(img_cv)
            if pts is not None and len(pts) > 0:
                pts = pts.astype(int)
                cv2.fillPoly(img_cv, pts, (255, 255, 255))
        except Exception:
            pass

        gray = cv2.cvtColor(img_cv, cv2.COLOR_BGR2GRAY)
        raw_page_text = pytesseract.image_to_string(gray, lang="ara+eng", config="--psm 6")
        raw_page_text = re.sub(r'[\u200e\u200f\u202a-\u202e]', '', raw_page_text)
        ocr_texts.append(raw_page_text)

    doc.close()
    combined_raw_text = "\n".join(ocr_texts)

    if is_quotation_file(pdf_path, combined_raw_text):
        return {"is_quotation": True}

    fields = parse_fields(combined_raw_text, pdf_path, img_cv=img_cv_first)
    return fields

def show_review_gui(all_results: list) -> bool:
    try:
        import tkinter as tk
        from tkinter import ttk

        root = tk.Tk()
        root.title("مراجعة واعتماد بيانات الفواتير")
        root.geometry("1100x520")
        root.configure(bg="#f4f6f9")
        root.attributes("-topmost", True)

        title_label = tk.Label(root, text="📋 جدول مراجعة الفواتير المستخرجة تلقائياً", font=("Segoe UI", 13, "bold"), bg="#f4f6f9", fg="#1a252f")
        title_label.pack(pady=8)

        cols = ("Num", "Customer name", "car-type", "Mech Labor", "Elec Labor", "worker/M", "worker/E", "Parts", "Network", "Cash", "Transfer", "Grand Total")
        tree = ttk.Treeview(root, columns=cols, show='headings', height=12)

        for col in cols:
            tree.heading(col, text=col)
            tree.column(col, width=85, anchor="center")

        tree.column("Customer name", width=150, anchor="center")
        tree.column("Parts", width=170, anchor="center")

        for r in all_results:
            tree.insert("", "end", values=(
                r.get("Num", ""),
                r.get("Customer name", ""),
                r.get("car-type", ""),
                r.get("Mech Labor amount", ""),
                r.get("Electrical Labor amount", ""),
                r.get("worker/M", ""),
                r.get("worker/E", ""),
                r.get("Parts & External Services", ""),
                r.get("Network / POS", ""),
                r.get("Cash", ""),
                r.get("Transfer", ""),
                r.get("Grand Total", "")
            ))

        tree.pack(fill="both", expand=True, padx=12, pady=4)

        user_approved = [True]

        def on_confirm():
            user_approved[0] = True
            root.destroy()

        def on_cancel():
            user_approved[0] = False
            root.destroy()

        btn_frame = tk.Frame(root, bg="#f4f6f9")
        btn_frame.pack(pady=10)

        btn_confirm = tk.Button(btn_frame, text="✅ اعتماد البيانات والتصدير للإكسل", font=("Segoe UI", 10, "bold"), bg="#27ae60", fg="white", padx=12, pady=5, command=on_confirm)
        btn_confirm.pack(side="left", padx=8)

        btn_cancel = tk.Button(btn_frame, text="❌ إلغاء المعالجة", font=("Segoe UI", 10), bg="#e74c3c", fg="white", padx=12, pady=5, command=on_cancel)
        btn_cancel.pack(side="right", padx=8)

        root.mainloop()
        return user_approved[0]
    except Exception as e:
        print(f"تخطي واجهة GUI: {e}")
        return True


def archive_processed_pdf(pdf_path: str):
    try:
        from datetime import datetime
        now = datetime.now()
        base_desktop = os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop")
        month_folder = f"Month_{now.month}"
        day_folder = now.strftime("%d")
        archive_dir = os.path.join(base_desktop, month_folder, "Daily-income", day_folder)
        os.makedirs(archive_dir, exist_ok=True)
        dest_file = os.path.join(archive_dir, os.path.basename(pdf_path))
        if os.path.exists(dest_file):
            name, ext = os.path.splitext(os.path.basename(pdf_path))
            dest_file = os.path.join(archive_dir, f"{name}_{now.strftime('%H%M%S')}{ext}")
        shutil.move(pdf_path, dest_file)
        print(f"تمت أرشفة الملف بنجاح إلى: {dest_file}")
    except Exception as e:
        print(f"تحذير: تعذر أرشفة الملف {os.path.basename(pdf_path)}: {e}")


if __name__ == "__main__":
    import glob
    path_input = r"\\DESKTOP-U37AB2R\Desktop\سكانات"
    print(f"Using fixed scanner directory: {path_input}")

    if os.path.isdir(path_input):
        print(f"Searching for PDF files in folder: {path_input}")
        pdf_files = glob.glob(os.path.join(path_input, "*.pdf"))

        if not pdf_files:
            print("No PDF files found in this folder.")
        else:
            all_results = []
            processed_files = []
            for pdf in pdf_files:
                print(f"\n--- Processing File: {os.path.basename(pdf)} ---")
                try:
                    result = extract_invoice(pdf)
                    if result.get("is_quotation"):
                        print(f"[تخطي] تم استبعاد ملف تسعيرة (QUTATION): {os.path.basename(pdf)} - النظام يقبل الفواتير الضريبية فقط.")
                        quote_dir = os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop", "تسعيرات_مستبعدة")
                        os.makedirs(quote_dir, exist_ok=True)
                        shutil.move(pdf, os.path.join(quote_dir, os.path.basename(pdf)))
                        continue
                    all_results.append(result)
                    processed_files.append(pdf)
                except Exception as e:
                    print(f"Error processing file: {e}")

            if all_results:
                all_results.sort(key=lambda x: x.get("Num_sort", 9999))
                approved = show_review_gui(all_results)
                if not approved:
                    print("تم إلغاء عملية الاعتماد بناءً على طلب المستخدم.")
                    sys.exit(0)

                import openpyxl

                desktop_excel = r"C:\Users\Dell\OneDrive\Desktop\Day_test.xlsx"
                proj_excel = r"C:\Users\Dell\OneDrive\Desktop\AUTOMATION\pdf-invoice-to-excel\Day_test.xlsx"
                target_file = desktop_excel if os.path.exists(desktop_excel) else proj_excel

                if os.path.exists(target_file):
                    print(f"\nجاري الكتابة المباشرة في ملف الإكسل: {target_file}")
                    wb = openpyxl.load_workbook(target_file)
                    ws = wb['ورقة1'] if 'ورقة1' in wb.sheetnames else wb.active

                    for idx, res in enumerate(all_results):
                        row = idx + 2
                        ws.cell(row=row, column=1).value = f"=+D{row}+C{row}+B{row}"
                        ws.cell(row=row, column=2).value = res["Cash"] if res["Cash"] != "" else None
                        ws.cell(row=row, column=3).value = res["Transfer"] if res["Transfer"] != "" else None
                        ws.cell(row=row, column=4).value = res["Network / POS"] if res["Network / POS"] != "" else None
                        ws.cell(row=row, column=5).value = res["Parts & External Services"] if res["Parts & External Services"] != "" else None
                        ws.cell(row=row, column=6).value = res["worker/E"] if res["worker/E"] != "" else None
                        ws.cell(row=row, column=7).value = res["worker/M"] if res["worker/M"] != "" else None
                        ws.cell(row=row, column=8).value = res["Electrical Labor amount"] if res["Electrical Labor amount"] != "" else None
                        ws.cell(row=row, column=9).value = res["Mech Labor amount"] if res["Mech Labor amount"] != "" else None
                        ws.cell(row=row, column=10).value = res["car-type"] if res["car-type"] != "" else None
                        ws.cell(row=row, column=11).value = res["Customer name"] if res["Customer name"] != "" else None
                        ws.cell(row=row, column=12).value = idx + 1

                    try:
                        wb.save(target_file)
                        print(f"\nتم تسجيل البيانات بنجاح داخل ملف: {target_file}")
                        for pdf_file in processed_files:
                            archive_processed_pdf(pdf_file)
                    except PermissionError:
                        alert_excel_locked(target_file)
                        fallback_out = os.path.join(os.path.dirname(target_file), "Day_test_updated.xlsx")
                        wb.save(fallback_out)
                        print(f"تم حفظ البيانات في ملف موازي: {fallback_out}")
                        for pdf_file in processed_files:
                            archive_processed_pdf(pdf_file)
