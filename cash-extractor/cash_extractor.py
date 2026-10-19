"""
Cash Income & Deposit Extractor
استخراج بيانات الكاش والإيداع من ملفات PDF وتصديرها إلى Excel
"""

import os
import re
import glob
import pdfplumber
import pandas as pd
from pathlib import Path
from datetime import datetime


# ============================================================
# إعدادات المسارات - عدّلها حسب موقع ملفاتك
# ============================================================
# طلب رقم الشهر من المستخدم
month_num = input("Enter the number of the month you want to Cash account (example, 9):").strip()

# بناء المسارات بناءً على الشهر المدخل
BASE_DIR = rf"C:\Users\Dell\OneDrive\Desktop\Month_{month_num}\Daily-income"
OUTPUT_DIR = rf"C:\Users\Dell\OneDrive\Desktop\Month_{month_num}\Cash-and-card\cash"
OUTPUT_EXCEL = os.path.join(
    OUTPUT_DIR, f"CASH-Income-And-Deposit-{month_num}.xlsx"
)

# أرقام الأيام المراد معالجتها
DAY_NUMBERS = list(range(1, 32))  # من 1 إلى 31


# ============================================================
# دوال استخراج البيانات
# ============================================================

def extract_text_from_pdf(pdf_path):
    """استخراج النص الكامل من ملف PDF"""
    full_text = ""
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                # النص من اليمين لليسار (نظراً لأن الملف عربي)
                text = page.extract_text()
                if text:
                    full_text += text + "\n"
        return full_text
    except Exception as e:
        print(f"❌ خطأ في قراءة {pdf_path}: {e}")
        return ""


def reverse_arabic_text(text):
    """عكس النص العربي وإصلاح ترتيب الحروف فيه ليُقرأ بشكل صحيح"""
    lines = text.split('\n')
    reversed_lines = []
    
    def rev_arabic(match):
        return match.group(0)[::-1]
        
    for line in lines:
        words = line.split()
        reversed_words = ' '.join(reversed(words))
        # عكس الحروف للكلمات العربية بعد عكس الكلمات
        fixed_line = re.sub(r'[\u0600-\u06FF]+', rev_arabic, reversed_words)
        reversed_lines.append(fixed_line)
        
    return '\n'.join(reversed_lines)


def sum_numbers_in_line(line):
    """حساب مجموع الأرقام المرتبطة بعلامة = في السطر"""
    parts = re.split(r'[-–—,\n]+', line)
    total = 0.0
    for p in parts:
        if '=' in p:
            nums = re.findall(r'(\d+(?:\.\d+)?)', p)
            if nums:
                total += float(nums[0])
    return total


def parse_expense_line(raw_line, cash_amount):
    """
    تفكيك سطر المصروفات إلى بنود صرف مستقلة إذا كان يحتوي على عدة بنود (مثل: 50=خالد - 50=نادر)
    أو إرجاعه كبند واحد إذا كان مصروفاً مفرداً.
    """
    if not raw_line or not raw_line.strip():
        return [{'amount': cash_amount, 'description': 'مصروف كاش'}]
        
    rev_line = reverse_arabic_text(raw_line).strip()
    
    # إذا لم يكن هناك '=' في النص، فهو بند مفرد
    if '=' not in raw_line:
        desc = rev_line.replace('\n', ' ')
        desc = re.sub(r'\s+', ' ', desc).strip()
        desc = re.sub(r'عبدا[اإ]لله', 'عبدالاله', desc)
        return [{'amount': cash_amount, 'description': desc}]
        
    # تفكيك البنود المفصولة بـ '-' أو ',' أو أسطر جديدة
    parts = re.split(r'[-–—,\n]+', raw_line)
    parts = [p.strip() for p in parts if p.strip()]
    
    sub_items = []
    # النص العربي من اليمين لليسار، لذا ترتيب الأجزاء يكون معكوساً
    for part in reversed(parts):
        if '=' not in part:
            continue
        sides = part.split('=', 1)
        side0 = sides[0].strip()
        side1 = sides[1].strip()
        
        m0 = re.match(r'^(\d+(?:\.\d+)?)$', side0)
        m1 = re.match(r'^(\d+(?:\.\d+)?)$', side1)
        
        if m0:
            amt = float(m0.group(1))
            desc = reverse_arabic_text(side1).strip()
        elif m1:
            amt = float(m1.group(1))
            desc = reverse_arabic_text(side0).strip()
        else:
            if re.match(r'^\d+(\.\d+)?$', re.sub(r'[^\d.]', '', side0)):
                amt = float(re.search(r'\d+(\.\d+)?', side0).group(0))
                desc = reverse_arabic_text(side1).strip()
            elif re.match(r'^\d+(\.\d+)?$', re.sub(r'[^\d.]', '', side1)):
                amt = float(re.search(r'\d+(\.\d+)?', side1).group(0))
                desc = reverse_arabic_text(side0).strip()
            else:
                num0 = re.search(r'(\d+(?:\.\d+)?)', side0)
                num1 = re.search(r'(\d+(?:\.\d+)?)', side1)
                if num0 and not num1:
                    amt = float(num0.group(1))
                    desc = reverse_arabic_text(side1).strip()
                elif num1 and not num0:
                    amt = float(num1.group(1))
                    desc = reverse_arabic_text(side0).strip()
                else:
                    amt = 0.0
                    desc = reverse_arabic_text(part).strip()
                    
        desc = re.sub(r'[\(\)\"\']', '', desc).strip()
        desc = re.sub(r'^(?:عهدة|عهده)\s+', '', desc)
        desc = re.sub(r'\s+(?:عهدة|عهده)$', '', desc)
        desc = re.sub(r'عبدا[اإ]لله', 'عبدالاله', desc)
        desc = re.sub(r'\s+', ' ', desc).strip()
        
        if amt > 0:
            sub_items.append({'amount': amt, 'description': desc})
            
    if not sub_items:
        desc = rev_line.replace('\n', ' ')
        desc = re.sub(r'عبدا[اإ]لله', 'عبدالاله', desc)
        return [{'amount': cash_amount, 'description': desc.strip()}]
        
    return sub_items


def extract_daily_cash_data(pdf_path, day_number):
    """استخراج بيانات الكاش والمصروفات ليوم معين بدقة وتفصيل كاملين"""
    
    text = extract_text_from_pdf(pdf_path)
    if not text:
        return None
    
    rev_text = reverse_arabic_text(text)
    
    data = {
        'day': day_number,
        'date': None,
        'invoices': [],
        'total_sales_cash': 0.0,
        'total_deposits': 0.0,
        'cash_remaining': 0.0,
        'cash_transfer_out': 0.0,
        'bank_transfer_out': 0.0,
        'expenses': [],
        'cash_deposit': 0.0
    }
    
    # استخراج التاريخ
    date_match = re.search(r'(\d{4}-\d{2}-\d{2})', rev_text)
    if date_match:
        data['date'] = date_match.group(1)
        
    # استخراج الفواتير والمصروفات عبر الجداول
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                for table in tables:
                    is_expense_section = False
                    header_row_idx = None
                    cash_col = 4
                    num_col = None
                    auth_col = None
                    
                    for r_idx, row in enumerate(table):
                        cleaned = [str(c).strip() if c is not None else '' for c in row]
                        joined = ' '.join(cleaned)
                        
                        if not is_expense_section:
                            # التعرف على ترويسة قسم المصروفات
                            if ('فيراصملا' in joined or 'المصاريف' in joined) and ('شاك' in joined or 'كاش' in joined or 'دنب' in joined or 'بند' in joined):
                                is_expense_section = True
                                header_row_idx = r_idx
                                for c_i, c in enumerate(cleaned):
                                    if 'شاك' in c or 'كاش' in c:
                                        cash_col = c_i
                                    elif 'ددع' in c or 'عدد' in c:
                                        num_col = c_i
                                    elif 'ديمعت' in c or 'تعميد' in c:
                                        auth_col = c_i
                                continue
                                
                            # 1. قراءة الإيرادات (قبل قسم المصروفات)
                            if len(row) >= 12:
                                num_str = cleaned[11]
                                if num_str.isdigit():
                                    cash_val_str = cleaned[1].replace(',', '')
                                    if cash_val_str and cash_val_str.lower() != 'none':
                                        try:
                                            amount = float(cash_val_str)
                                            if amount > 0:
                                                name_str = cleaned[10]
                                                reversed_name = reverse_arabic_text(name_str).replace('\n', ' ')
                                                invoice_num = ""
                                                customer_name = reversed_name
                                                match_num = re.search(r'(\d{3,})', reversed_name)
                                                if match_num:
                                                    invoice_num = match_num.group(1)
                                                    customer_name = reversed_name.replace(invoice_num, '').strip()
                                                    
                                                data['invoices'].append({
                                                    'invoice_number': invoice_num,
                                                    'customer_name': customer_name,
                                                    'amount1': amount,
                                                    'amount2': amount
                                                })
                                        except ValueError:
                                            pass
                                            
                    # 2. قراءة وتفكيك المصروفات الكاش (داخل قسم المصروفات)
                    if is_expense_section and header_row_idx is not None:
                        raw_expense_rows = []
                        for r_idx in range(header_row_idx + 1, len(table)):
                            row = table[r_idx]
                            cleaned = [str(c).strip() if c is not None else '' for c in row]
                            joined = ' '.join(cleaned)
                            if any(term in joined for term in ['تاداريا', 'ايرادات', 'تافوصرملا ليامجا', 'اجمالي المصروفات', 'يقبتملا', 'المتبقي']):
                                break
                                
                            cash_val = 0.0
                            if cash_col < len(cleaned):
                                cash_str = cleaned[cash_col].replace(',', '')
                                try:
                                    cash_val = float(cash_str)
                                except ValueError:
                                    cash_val = 0.0
                                    
                            end_col = num_col if (num_col and num_col > cash_col) else len(cleaned) - 1
                            desc_parts = [c for c in cleaned[cash_col+1:end_col] if c]
                            raw_desc = '\n'.join(desc_parts)
                            
                            auth_name = ""
                            if auth_col and auth_col < len(cleaned):
                                auth_name = reverse_arabic_text(cleaned[auth_col]).strip()
                            elif cleaned[-1]:
                                auth_name = reverse_arabic_text(cleaned[-1]).strip()
                                
                            raw_expense_rows.append({
                                'cash': cash_val,
                                'raw_desc': raw_desc,
                                'auth_name': auth_name
                            })
                            
                        # مطابقة وتوزيع الأسطر للبيانات متعددة السطور
                        for i, r in enumerate(raw_expense_rows):
                            if r['cash'] > 0:
                                lines = [l.strip() for l in r['raw_desc'].split('\n') if l.strip()]
                                if len(lines) > 1:
                                    matched_line = None
                                    for l in lines:
                                        if abs(sum_numbers_in_line(l) - r['cash']) < 0.01:
                                            matched_line = l
                                            break
                                    if matched_line:
                                        r['raw_desc'] = matched_line
                                    else:
                                        r['raw_desc'] = lines[0]
                                elif not r['raw_desc']:
                                    # إذا كان البيان فارغاً، نبحث في الصفوف السابقة عن سطر متبقي يطابق المبلغ
                                    for prev in reversed(raw_expense_rows[:i]):
                                        if prev['raw_desc']:
                                            p_lines = [l.strip() for l in prev['raw_desc'].split('\n') if l.strip()]
                                            matched_line = None
                                            for l in p_lines:
                                                if abs(sum_numbers_in_line(l) - r['cash']) < 0.01:
                                                    matched_line = l
                                                    break
                                            if matched_line:
                                                r['raw_desc'] = matched_line
                                                break
                                            elif len(p_lines) > 1:
                                                r['raw_desc'] = p_lines[-1]
                                                break
                                    if not r['raw_desc'] and r['auth_name']:
                                        r['raw_desc'] = r['auth_name']
                                        
                        # استخراج الإيداعات النقدية (مثل: 7000 = ايداع نقدي)
                        for r in raw_expense_rows:
                            raw_d = r['raw_desc']
                            rev_d = reverse_arabic_text(raw_d)
                            if 'يدقن عاديا' in raw_d or 'ايداع نقدي' in rev_d or 'إيداع نقدي' in rev_d or ('ايداع' in rev_d and 'نقدي' in rev_d):
                                nums = re.findall(r'(\d+(?:\.\d+)?)', raw_d)
                                if nums:
                                    data['cash_deposit'] += float(nums[0])

                        # تفكيك بنود الصرف وإضافتها كبنود مستقلة
                        for r in raw_expense_rows:
                            if r['cash'] > 0:
                                items = parse_expense_line(r['raw_desc'], r['cash'])
                                data['expenses'].extend(items)
                                
    except Exception as e:
        print(f"❌ خطأ في استخراج الجداول من ملف الـ PDF: {e}")

    # حساب الإجماليات بدقة من الفواتير والمصروفات المستخرجة
    data['total_sales_cash'] = sum(inv['amount1'] for inv in data['invoices'])
    data['total_deposits'] = sum(exp['amount'] for exp in data['expenses'])
    data['cash_remaining'] = data['total_sales_cash'] - data['total_deposits'] - data['cash_deposit']
    
    return data


def find_pdf_file(day_number):
    """البحث عن ملف PDF الخاص بتقرير اليوم في مجلد اليوم"""
    
    # البحث عن المجلد بصيغة الرقم العادي (مثل 2) أو بصيغة رقمين (مثل 02)
    day_folder = os.path.join(BASE_DIR, str(day_number))
    day_folder_0 = os.path.join(BASE_DIR, f"{day_number:02d}")
    
    if os.path.exists(day_folder_0):
        day_folder = day_folder_0
    elif not os.path.exists(day_folder):
        return None
    
    # 1. البحث أولاً عن ملف التقرير اليومي الشامل مثل Day-1.pdf أو Day 1.pdf
    priority_patterns = [
        os.path.join(day_folder, f"Day-{day_number}.pdf"),
        os.path.join(day_folder, f"Day {day_number}.pdf"),
        os.path.join(day_folder, f"Day_{day_number}.pdf"),
        os.path.join(day_folder, f"day-{day_number}.pdf"),
        os.path.join(day_folder, f"day {day_number}.pdf"),
        os.path.join(day_folder, f"*{day_number}*.pdf"),
    ]
    
    for pattern in priority_patterns:
        matches = glob.glob(pattern)
        if matches:
            return matches[0]
            
    # 2. إذا لم نجد ملفاً بهذا الاسم، نأخذ أي ملف PDF موجود في المجلد
    pdf_files = glob.glob(os.path.join(day_folder, "*.pdf"))
    if pdf_files:
        return pdf_files[0]
    return None


def process_all_days():
    """معالجة جميع الأيام وتجميع البيانات"""
    
    all_data = []
    summary = {
        'total_sales_cash': 0.0,
        'total_expenses': 0.0,
        'total_cash_deposits': 0.0,
        'total_cash_remaining': 0.0,
    }
    
    print("=" * 60)
    print("🚀 بدء معالجة ملفات الكاش الشهرية")
    print("=" * 60)
    
    for day_num in DAY_NUMBERS:
        pdf_path = find_pdf_file(day_num)
        
        if not pdf_path:
            print(f"⚠️  اليوم {day_num}: لم يتم العثور على ملف")
            continue
        
        print(f"📄 اليوم {day_num}: {os.path.basename(pdf_path)}")
        
        data = extract_daily_cash_data(pdf_path, day_num)
        if data:
            all_data.append(data)
            summary['total_sales_cash'] += data['total_sales_cash']
            summary['total_expenses'] += data['total_deposits']
            summary['total_cash_deposits'] += data['cash_deposit']
            summary['total_cash_remaining'] += data['cash_remaining']
    
    return all_data, summary


def create_excel_report(all_data, summary):
    """إنشاء ملف Excel موحد بنفس التنسيق المطلوب في الصورة"""
    
    output_path = OUTPUT_EXCEL
    
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "التقرير"
    # تفعيل اتجاه الورقة من اليمين لليسار
    ws.sheet_view.rightToLeft = True

    # تعريف التنسيقات
    header_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    gray_fill = PatternFill(start_color="A6A6A6", end_color="A6A6A6", fill_type="solid")
    
    green_fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
    red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    
    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # كتابة العناوين (من اليمين لليسار)
    headers = ['التاريخ', 'رقم الفاتورة', 'استلام مبلغ اجور', 'صرف', 'بيان الصرف', 'ايداعات']
    ws.append(headers)
    
    for col_idx in range(1, 7):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = border
        
    # تحديد عرض الأعمدة
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 15

    total_income = 0
    total_expense = 0
    total_cash_deposits = 0

    current_row = 2
    for day_idx, day_data in enumerate(all_data):
        date_str = day_data['date'] if day_data['date'] else f"اليوم {day_data['day']}"
        invoices = day_data['invoices']
        expenses = day_data['expenses']
        deposit_val = day_data.get('cash_deposit', 0.0)
        
        max_len = max(len(invoices), len(expenses), 1 if deposit_val > 0 else 0)
        if max_len == 0:
            continue
            
        for i in range(max_len):
            row_vals = [""] * 6
            
            # عمود A: التاريخ (يظهر مع كل فاتورة، أو في أول سطر لليوم إذا لم تكن هناك فواتير)
            if i < len(invoices):
                row_vals[0] = date_str
            elif len(invoices) == 0 and i == 0:
                row_vals[0] = date_str
            
            # الفواتير: رقم الفاتورة (B) والمبلغ (C)
            if i < len(invoices):
                row_vals[1] = invoices[i]['invoice_number']
                row_vals[2] = invoices[i]['amount1']
                total_income += invoices[i]['amount1']
                
            # المصروفات: المبلغ (D) والبيان (E)
            if i < len(expenses):
                row_vals[3] = expenses[i]['amount']
                row_vals[4] = expenses[i]['description']
                total_expense += expenses[i]['amount']
                
            # عمود F: إيداعات نقدية (تسجيل المبلغ في نفس صف تاريخ اليوم)
            if i == 0 and deposit_val > 0:
                row_vals[5] = deposit_val
                total_cash_deposits += deposit_val
            
            for col_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.alignment = center_align
                cell.border = border
                if col_idx == 3 and val != "":
                    cell.font = Font(color="00B050", bold=True)
                elif col_idx == 6 and val != "":
                    cell.font = bold_font
                
            current_row += 1
            
        # إضافة صف فاصل رمادي بين الأيام
        if day_idx < len(all_data) - 1:
            for col_idx in range(1, 7):
                cell = ws.cell(row=current_row, column=col_idx, value="")
                cell.fill = gray_fill
                cell.border = border
            current_row += 1

    # إضافة صف الإجمالي بأسفل الجدول مباشرة
    ws.cell(row=current_row, column=1, value="الاجمالي").font = bold_font
    cell_tot_inc = ws.cell(row=current_row, column=3, value=total_income)
    cell_tot_inc.font = Font(color="00B050", bold=True)
    ws.cell(row=current_row, column=4, value=total_expense).font = bold_font
    if total_cash_deposits > 0:
        ws.cell(row=current_row, column=6, value=total_cash_deposits).font = bold_font
    
    for col_idx in range(1, 7):
        cell = ws.cell(row=current_row, column=col_idx)
        cell.alignment = center_align
        cell.border = border
    
    current_row += 2
    
    # مربع الملخص الملون (أسفل الجدول)
    total_income = round(total_income, 2)
    total_expense = round(total_expense, 2)
    cash_deposit = round(total_cash_deposits, 2)
    cash_remaining = round(total_income - total_expense - cash_deposit, 2)
    
    summary_data = [
        ("اجمالي استلام", total_income, green_fill),
        ("اجمالي الصرف", total_expense, red_fill),
        ("ايداع النقدي", cash_deposit, yellow_fill),
        ("كاش المتبقي", cash_remaining, None),
    ]
    
    for desc, val, fill_color in summary_data:
        cell_desc = ws.cell(row=current_row, column=1, value=desc)
        cell_val = ws.cell(row=current_row, column=2, value=val)
        
        cell_desc.font = bold_font
        cell_val.font = bold_font
        
        cell_desc.alignment = center_align
        cell_val.alignment = center_align
        
        if fill_color:
            cell_desc.fill = fill_color
            cell_val.fill = fill_color
            
        current_row += 1

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    
    print(f"\n✅ تم إنشاء الملف بنجاح بالشكل المخصص: {output_path}")
    return output_path


# ============================================================
# الدالة الرئيسية
# ============================================================
def main():
    start_time = datetime.now()
    
    # معالجة جميع الأيام
    all_data, summary = process_all_days()
    
    # إنشاء ملف Excel
    if all_data:
        output_file = create_excel_report(all_data, summary)
        
        print("\n" + "=" * 60)
        print("📊 الملخص النهائي")
        print("=" * 60)
        print(f"عدد الأيام المعالجة : {len(all_data)}")
        print(f"إجمالي المبيعات     : {summary['total_sales_cash']:.2f} ريال")
        print(f"إجمالي المصروفات    : {summary['total_expenses']:.2f} ريال")
        print(f"إجمالي الإيداعات     : {summary['total_cash_deposits']:.2f} ريال")
        print(f"الكاش المتبقي       : {summary['total_cash_remaining']:.2f} ريال")
        print(f"الملف المُنشأ        : {output_file}")
        print(f"الوقت المستغرق      : {datetime.now() - start_time}")
        print("=" * 60)
    else:
        print("❌ لم يتم العثور على أي بيانات")


if __name__ == "__main__":
    main()
