import io
import re
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(page_title='تسويات نقاط البيع', layout='wide')
st.title('تقرير إيداعات وخصومات تسويات نقاط البيع')
st.caption('XZ = إيداع، FZ = خصم رسوم | التصفية حسب تاريخ التسوية، مع حفظ تاريخ النزول البنكي')

FIELDS = ['صف المصدر','تاريخ نزول البنك','تاريخ التسوية','أيام التأخير','النوع','رمز البنك','رمز الشبكة','رقم التسوية','المرجع','المبلغ','مدين','دائن','التفاصيل','الوصف','حالة التدقيق','ملاحظات']

def money(value):
    if pd.isna(value) or str(value).strip() in ('','-'):
        return Decimal('0')
    try:
        return Decimal(str(value).replace(',','').replace('ر.س','').strip())
    except (InvalidOperation, ValueError):
        return None

def settlement(description):
    m = re.search(r'\b([A-Z]{2})(\d{6})\b', description.upper())
    if not m:
        return '', None
    try:
        return m.group(1), datetime.strptime(m.group(2), '%y%m%d').date()
    except ValueError:
        return m.group(1), None

def read_bank(data):
    raw = pd.read_excel(io.BytesIO(data), header=None)
    header = None
    for i, row in raw.iterrows():
        vals = [str(x).strip() for x in row.tolist()]
        if '[SA]Processing Date' in vals and 'مدين' in vals and 'دائن' in vals:
            header = i
            break
    if header is None:
        raise ValueError('لم أجد صف عناوين كشف البنك (تاريخ التنفيذ، مدين، دائن).')
    names = [str(x).strip() for x in raw.iloc[header].tolist()]
    debit_col, credit_col = names.index('مدين'), names.index('دائن')
    records = []
    for i in range(header + 1, len(raw)):
        row = raw.iloc[i]
        bank = pd.to_datetime(row.iloc[0], errors='coerce')
        if pd.isna(bank):
            continue
        code = str(row.iloc[3]).strip().upper()
        if code not in ('XZ','FZ'):
            continue
        desc = str(row.iloc[2]) if pd.notna(row.iloc[2]) else ''
        detail = str(row.iloc[1]) if pd.notna(row.iloc[1]) else ''
        debit, credit = money(row.iloc[debit_col]), money(row.iloc[credit_col])
        network, settled = settlement(desc + ' ' + detail)
        recon = re.search(r'RECON\s*NO\s*:\s*(\d+)', desc + ' ' + detail, re.I)
        ref = re.search(r'REFERENCE\s*:\s*(\d+)', desc + ' ' + detail, re.I)
        notes = []
        if debit is None or credit is None:
            notes.append('مبلغ غير قابل للقراءة')
        debit = debit if debit is not None else Decimal('0')
        credit = credit if credit is not None else Decimal('0')
        direction_ok = (code == 'XZ' and credit > 0 and debit == 0) or (code == 'FZ' and debit > 0 and credit == 0)
        if not direction_ok:
            notes.append('اتجاه المبلغ غير صحيح أو المبلغ صفر')
        if settled is None:
            notes.append('تعذر استخراج تاريخ التسوية')
        lag = (bank.date() - settled).days if settled else None
        if lag is not None and lag != 1:
            notes.append(f'فارق النزول {lag} يوم (قد يكون عطلة)')
        if not recon:
            notes.append('رقم التسوية غير موجود')
        if not ref:
            notes.append('المرجع غير موجود')
        records.append({'صف المصدر':i+1,'تاريخ نزول البنك':bank.date(),'تاريخ التسوية':settled,'أيام التأخير':lag,
                        'النوع':'إيداع نقاط البيع' if code=='XZ' else 'خصم رسوم نقاط البيع',
                        'رمز البنك':code,'رمز الشبكة':network,'رقم التسوية':recon.group(1) if recon else '',
                        'المرجع':ref.group(1) if ref else '', 'المبلغ':float(credit if code=='XZ' else debit),
                        'مدين':float(debit),'دائن':float(credit),'التفاصيل':detail,'الوصف':desc,
                        'حالة التدقيق':'محتسب' if direction_ok and settled else 'مستبعد',
                        'ملاحظات':'؛ '.join(notes)})
    df = pd.DataFrame(records, columns=FIELDS)
    if df.empty:
        raise ValueError('لا توجد حركات XZ أو FZ في كشف البنك.')
    keys = ['تاريخ نزول البنك','تاريخ التسوية','رمز البنك','رمز الشبكة','رقم التسوية','المرجع','المبلغ']
    duplicated = df.duplicated(subset=keys, keep=False)
    for i in df.index[duplicated]:
        note = df.at[i,'ملاحظات']
        df.at[i,'ملاحظات'] = (note + '؛ ' if note else '') + 'تكرار محتمل: راجع المرجع والتسوية'
    return df

def daily_table(frame, date_col, start=None, end=None):
    valid = frame[frame['حالة التدقيق']=='محتسب']
    if valid.empty:
        table = pd.DataFrame(columns=[date_col,'إجمالي الإيداعات','إجمالي الخصومات'])
    else:
        table = valid.pivot_table(index=date_col, columns='رمز البنك', values='المبلغ', aggfunc='sum', fill_value=0).reset_index()
        table.columns.name = None
        for code in ('XZ','FZ'):
            if code not in table.columns:
                table[code] = 0.0
        table = table.rename(columns={'XZ':'إجمالي الإيداعات','FZ':'إجمالي الخصومات'})
    if start is not None and end is not None:
        days = pd.date_range(start,end,freq='D').date
        table = table.set_index(date_col).reindex(days,fill_value=0).rename_axis(date_col).reset_index()
    table['صافي الإيداعات بعد الرسوم (للمقارنة)'] = table['إجمالي الإيداعات'] - table['إجمالي الخصومات']
    return table.sort_values(date_col)

def excel_report(frame, daily, bank_daily, start, end, excluded_unknown):
    valid = frame[frame['حالة التدقيق']=='محتسب']
    deposits = round(valid.loc[valid['رمز البنك']=='XZ','المبلغ'].sum(),2)
    fees = round(valid.loc[valid['رمز البنك']=='FZ','المبلغ'].sum(),2)
    summary = pd.DataFrame([
        ('بداية الفترة',start.isoformat()),('نهاية الفترة',end.isoformat()),
        ('إجمالي الإيداعات',deposits),('إجمالي الخصومات',fees),('الصافي للمقارنة',round(deposits-fees,2)),
        ('عدد الإيداعات',int((valid['رمز البنك']=='XZ').sum())),
        ('عدد الخصومات',int((valid['رمز البنك']=='FZ').sum())),
        ('حركات مستبعدة ضمن الفترة',int((frame['حالة التدقيق']!='محتسب').sum())),
        ('حركات بلا تاريخ تسوية في الملف الأصلي',excluded_unknown),
        ('حركات عليها ملاحظات',int(frame['ملاحظات'].ne('').sum()))], columns=['البيان','القيمة'])
    network = valid.pivot_table(index=['تاريخ التسوية','رمز الشبكة'],columns='رمز البنك',values='المبلغ',aggfunc='sum',fill_value=0).reset_index() if not valid.empty else pd.DataFrame(columns=['تاريخ التسوية','رمز الشبكة','XZ','FZ'])
    network.columns.name = None
    for code in ('XZ','FZ'):
        if code not in network.columns:
            network[code] = 0.0
    network = network.rename(columns={'XZ':'إجمالي الإيداعات','FZ':'إجمالي الخصومات'})
    sheets = [('الملخص',summary),('حسب تاريخ التسوية',daily),('حسب تاريخ البنك',bank_daily),('حسب الشبكة',network),
              ('كل الحركات',frame),('حركات للمراجعة',frame[frame['ملاحظات'].ne('')]),
              ('حركات مستبعدة',frame[frame['حالة التدقيق']!='محتسب'])]
    wb = Workbook()
    wb.remove(wb.active)
    for name, data in sheets:
        ws = wb.create_sheet(name)
        ws.sheet_view.rightToLeft = True
        ws.append(list(data.columns))
        for record in data.itertuples(index=False,name=None):
            ws.append([v.to_pydatetime() if isinstance(v,pd.Timestamp) else v for v in record])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill('solid',fgColor='17365D')
            cell.font = Font(color='FFFFFF',bold=True)
            cell.alignment = Alignment(horizontal='center')
        for j,col in enumerate(data.columns,1):
            ws.column_dimensions[get_column_letter(j)].width = min(65,max(18,len(str(col))+5))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool):
                    cell.number_format = '#,##0.00;[Red](#,##0.00)'
                elif isinstance(cell.value,(datetime,date)):
                    cell.number_format = 'yyyy-mm-dd'
                if name in ('حركات للمراجعة','حركات مستبعدة'):
                    cell.fill = PatternFill('solid',fgColor='FFF2CC')
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(),deposits,fees

uploaded = st.file_uploader('ارفع كشف البنك Excel',type=['xlsx'])
if uploaded is not None:
    try:
        all_df = read_bank(uploaded.getvalue())
        known = all_df['تاريخ التسوية'].dropna()
        if known.empty:
            st.error('لا توجد تواريخ تسوية صالحة في الكشف؛ راجع وصف الحركات.')
            st.stop()
        first,last = known.min(),known.max()
        st.subheader('فلتر فترة التسوية')
        mode = st.radio('عرض التقرير', ['جميع الأيام','فترة مخصصة'], horizontal=True)
        start,end = first,last
        if mode == 'فترة مخصصة':
            left,right = st.columns(2)
            with left:
                start = st.date_input('من تاريخ',value=first,min_value=first,max_value=last)
            with right:
                end = st.date_input('إلى تاريخ',value=last,min_value=first,max_value=last)
        if start > end:
            st.error('تاريخ البداية أكبر من تاريخ النهاية.')
            st.stop()
        selected = all_df[all_df['تاريخ التسوية'].apply(lambda d: d is not None and start <= d <= end)].copy()
        unknown = int(all_df['تاريخ التسوية'].isna().sum())
        daily = daily_table(selected,'تاريخ التسوية',start,end)
        bank_daily = daily_table(selected,'تاريخ نزول البنك')
        result,deposits,fees = excel_report(selected,daily,bank_daily,start,end,unknown)
        st.caption(f'من {start:%Y-%m-%d} إلى {end:%Y-%m-%d} — حسب تاريخ التسوية')
        a,b,c,d = st.columns(4)
        a.metric('إجمالي الإيداعات',f'{deposits:,.2f} ر.س')
        b.metric('إجمالي الخصومات',f'{fees:,.2f} ر.س')
        c.metric('الصافي للمقارنة',f'{deposits-fees:,.2f} ر.س')
        d.metric('حركات تحتاج مراجعة',int(selected['ملاحظات'].ne('').sum())+unknown)
        if unknown:
            st.warning(f'هناك {unknown} حركة بلا تاريخ تسوية؛ لم تدخل في إجمالي الفترة المحددة. راجع الملف الأصلي.')
        st.subheader('التسويات اليومية (تشمل الأيام بدون حركات)')
        st.dataframe(daily,width='stretch',hide_index=True)
        st.subheader('تفاصيل الحركات')
        st.dataframe(selected,width='stretch',hide_index=True)
        st.download_button('تنزيل تقرير Excel',result,file_name=f'تقرير_التسويات_{start:%Y%m%d}_{end:%Y%m%d}.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as exc:
        st.error(f'تعذر معالجة الكشف: {exc}')
