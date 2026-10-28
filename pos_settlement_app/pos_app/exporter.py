"""
وحدة توليد وتصدير جداول وتقارير Excel لتسويات نقاط البيع
"""

import io
from datetime import datetime, date
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from constants import (
    CODE_DEPOSIT,
    CODE_FEE,
    AUDIT_STATUS_VALID
)

# ألوان التنسيق القياسية
COLOR_HEADER_BG = '17365D'       # كحلي رسمي
COLOR_HEADER_FONT = 'FFFFFF'     # أبيض
COLOR_WARNING_BG = 'FFF2CC'      # أصفر تنبيهي خفيف


def daily_table(frame, date_col, start=None, end=None):
    """
    توليد جدول التسويات اليومية (إيداعات، خصومات، صافي) مع ملء الأيام الفارغة إن حُدد النطاق.
    """
    valid = frame[frame['حالة التدقيق'] == AUDIT_STATUS_VALID]

    if valid.empty:
        table = pd.DataFrame(columns=[date_col, 'إجمالي الإيداعات', 'إجمالي الخصومات'])
    else:
        table = valid.pivot_table(
            index=date_col,
            columns='رمز البنك',
            values='المبلغ',
            aggfunc='sum',
            fill_value=0.0
        ).reset_index()
        table.columns.name = None

        for code in (CODE_DEPOSIT, CODE_FEE):
            if code not in table.columns:
                table[code] = 0.0

        table = table.rename(columns={
            CODE_DEPOSIT: 'إجمالي الإيداعات',
            CODE_FEE: 'إجمالي الخصومات'
        })

    if start is not None and end is not None:
        days = pd.date_range(start, end, freq='D').date
        table = table.set_index(date_col).reindex(days, fill_value=0.0).rename_axis(date_col).reset_index()

    # احتساب الصافي
    table['صافي الإيداعات بعد الرسوم (للمقارنة)'] = (
        table['إجمالي الإيداعات'].astype(float) - table['إجمالي الخصومات'].astype(float)
    ).round(2)

    return table.sort_values(date_col)


def _apply_sheet_styling(ws, data_df, is_warning_sheet=False):
    """
    تطبيق التنسيقات الاحترافية على ورقة Excel: RTL، تجميد الأسطر، الفلاتر،
    الألوان، وحساب عروض الأعمدة بناءً على المحتوى الفعلي.
    """
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    # تنسيق رأس الجدول
    header_fill = PatternFill('solid', fgColor=COLOR_HEADER_BG)
    header_font = Font(name='Segoe UI', size=11, color=COLOR_HEADER_FONT, bold=True)
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    warning_fill = PatternFill('solid', fgColor=COLOR_WARNING_BG) if is_warning_sheet else None

    # تنسيق الخلايا وتحديد عرض الأعمدة ديناميكياً
    for col_idx, col_name in enumerate(data_df.columns, 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(col_name))

        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value

            if warning_fill:
                cell.fill = warning_fill

            if isinstance(val, (int, float)) and not isinstance(val, bool):
                cell.number_format = '#,##0.00;[Red](#,##0.00)'
                cell.alignment = Alignment(horizontal='right', vertical='center')
                val_len = len(f'{val:,.2f}')
            elif isinstance(val, (datetime, date)):
                cell.number_format = 'yyyy-mm-dd'
                cell.alignment = Alignment(horizontal='center', vertical='center')
                val_len = 10
            else:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                val_len = len(str(val)) if val is not None else 0

            if val_len > max_len:
                max_len = val_len

        # احتساب العرض مع حواف أمان كافية وبحد أقصى لمنع الاتساع المفرط
        ws.column_dimensions[col_letter].width = min(60, max(14, max_len + 4))

    ws.row_dimensions[1].height = 28


def excel_report(frame, daily, bank_daily, start, end, excluded_unknown):
    """
    توليد التقرير المالي الشامل بصيغة Excel متعدد التبويبات باللغة العربية.
    """
    valid = frame[frame['حالة التدقيق'] == AUDIT_STATUS_VALID]
    deposits = round(float(valid.loc[valid['رمز البنك'] == CODE_DEPOSIT, 'المبلغ'].sum()), 2)
    fees = round(float(valid.loc[valid['رمز البنك'] == CODE_FEE, 'المبلغ'].sum()), 2)
    net_val = round(deposits - fees, 2)

    summary = pd.DataFrame([
        ('بداية الفترة', start.isoformat()),
        ('نهاية الفترة', end.isoformat()),
        ('إجمالي الإيداعات', deposits),
        ('إجمالي الخصومات', fees),
        ('الصافي للمقارنة', net_val),
        ('عدد الإيداعات', int((valid['رمز البنك'] == CODE_DEPOSIT).sum())),
        ('عدد الخصومات', int((valid['رمز البنك'] == CODE_FEE).sum())),
        ('حركات مستبعدة ضمن الفترة', int((frame['حالة التدقيق'] != AUDIT_STATUS_VALID).sum())),
        ('حركات بلا تاريخ تسوية في الملف الأصلي', excluded_unknown),
        ('حركات عليها ملاحظات', int(frame['ملاحظات'].ne('').sum()))
    ], columns=['البيان', 'القيمة'])

    if not valid.empty:
        network = valid.pivot_table(
            index=['تاريخ التسوية', 'رمز الشبكة'],
            columns='رمز البنك',
            values='المبلغ',
            aggfunc='sum',
            fill_value=0.0
        ).reset_index()
        network.columns.name = None
        for code in (CODE_DEPOSIT, CODE_FEE):
            if code not in network.columns:
                network[code] = 0.0
        network = network.rename(columns={
            CODE_DEPOSIT: 'إجمالي الإيداعات',
            CODE_FEE: 'إجمالي الخصومات'
        })
    else:
        network = pd.DataFrame(columns=['تاريخ التسوية', 'رمز الشبكة', 'إجمالي الإيداعات', 'إجمالي الخصومات'])

    sheets = [
        ('الملخص', summary, False),
        ('حسب تاريخ التسوية', daily, False),
        ('حسب تاريخ البنك', bank_daily, False),
        ('حسب الشبكة', network, False),
        ('كل الحركات', frame, False),
        ('حركات للمراجعة', frame[frame['ملاحظات'].ne('')], True),
        ('حركات مستبعدة', frame[frame['حالة التدقيق'] != AUDIT_STATUS_VALID], True)
    ]

    wb = Workbook()
    wb.remove(wb.active)

    for name, data, is_warning in sheets:
        ws = wb.create_sheet(name)
        ws.append(list(data.columns))

        for record in data.itertuples(index=False, name=None):
            row_vals = [
                v.to_pydatetime() if isinstance(v, pd.Timestamp) else v
                for v in record
            ]
            ws.append(row_vals)

        _apply_sheet_styling(ws, data, is_warning_sheet=is_warning)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), deposits, fees
